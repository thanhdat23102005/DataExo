"""
Generate the Early-Stage Presentation deck (Week 8, 40 marks).

Built for a 10-minute slot. The presentation rubric scores Organization,
Subject Knowledge, Mechanics, Eye Contact and Elocution - the last two are
delivery, so every slide carries speaker notes with a timing and what to
say. Slides hold headlines, not paragraphs, so the speaker is looking at
the audience instead of reading the screen.

Usage:
    python3 build_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x39, 0x87, 0xE5)
TEAL = RGBColor(0x1E, 0x8A, 0x7A)
RED = RGBColor(0xC0, 0x39, 0x2B)
AMBER = RGBColor(0xC9, 0x7A, 0x16)
GREY = RGBColor(0x6B, 0x6B, 0x72)
DARK = RGBColor(0x2A, 0x2A, 0x30)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ASSETS = Path(__file__).parent / "assets"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, x, y, w, h, text, size=18, color=DARK, bold=False,
            align=PP_ALIGN.LEFT, spacing=1.0, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def band(slide, color=NAVY, height=Inches(0.14), y=Inches(0.0)):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), y, SLIDE_W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def slide_title(slide, title, kicker=None):
    band(slide)
    if kicker:
        textbox(slide, Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.35),
                kicker.upper(), size=12, color=BLUE, bold=True)
        y = Inches(0.78)
    else:
        y = Inches(0.55)
    textbox(slide, Inches(0.8), y, Inches(11.7), Inches(0.9),
            title, size=30, color=NAVY, bold=True)


def bullets(slide, items, x=Inches(0.9), y=Inches(1.9), w=Inches(11.5),
            size=19, gap=Inches(0.72)):
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            head, sub = item
        else:
            head, sub = item, None
        textbox(slide, x, y + gap * i, w, Inches(0.45), head, size=size,
                color=DARK, bold=True)
        if sub:
            textbox(slide, x, y + gap * i + Inches(0.34), w, Inches(0.35),
                    sub, size=14, color=GREY)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ---------------- 1. Title ----------------
    s = blank(prs)
    band(s, NAVY, Inches(2.6), Inches(0))
    textbox(s, Inches(0.9), Inches(0.75), Inches(11.5), Inches(0.9),
            "DataExodus", size=46, color=WHITE, bold=True)
    textbox(s, Inches(0.95), Inches(1.72), Inches(11.5), Inches(0.6),
            "Where does your data actually go?", size=20, color=RGBColor(0xBD, 0xD6, 0xF5))
    textbox(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.8),
            "Early-Stage Presentation  ·  Project Plan", size=22, color=NAVY, bold=True)
    textbox(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.6),
            ["Thanh Dat Phan  ·  [Student ID]",
             "[Unit code and name]  ·  Week 8",
             "github.com/thanhdat23102005/DataExo"], size=16, color=GREY, spacing=1.35)
    notes(s, """
[0:00-0:15]
Good morning. My project is called DataExodus, and it answers one question:
when you open an ordinary Australian website, where does your data actually go?
Over the next ten minutes I'll cover the problem, what I'm building, how I'm
building it, and what is already working.
Do not read the slide. Say the question, then pause.
""")

    # ---------------- 2. The problem ----------------
    s = blank(prs)
    slide_title(s, "A hash is not anonymous", kicker="the problem")
    textbox(s, Inches(0.9), Inches(1.85), Inches(11.5), Inches(0.9),
            "you@example.com   →   973dfe463ec857…", size=26, color=NAVY, bold=True)
    textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(0.6),
            "Companies call this anonymised. It is not.", size=20, color=RED, bold=True)
    bullets(s, [
        ("The same email always produces the same hash",
         "so that value is a stable name for you"),
        ("It follows you between unrelated websites",
         "any site receiving the same hash recognises the same person"),
        ("The raw address never crosses the wire",
         "which is exactly why the practice looks safe and is not"),
    ], y=Inches(3.5), size=19)
    textbox(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.5),
            "US Federal Trade Commission, 2024: \"Hashing is not anonymization\"",
            size=14, color=GREY)
    notes(s, """
[0:15-1:25]
This is the heart of the project. A tracker rarely sends your email address in
plain text - it sends the hash of it, and calls that anonymised.

But hashing is deterministic. The same email always produces the same hash. So
that value is still a name for you, it just isn't readable. Any site that
receives it recognises the same person.

The FTC said exactly this in 2024. My project doesn't argue the point - it
demonstrates it, on the audience's own traffic.

Pause after "It is not." Let it land.
""")

    # ---------------- 3. Background ----------------
    s = blank(prs)
    slide_title(s, "Why this matters in Australia", kicker="background")
    bullets(s, [
        ("APP 8 — Privacy Act 1988",
         "organisations disclosing personal information overseas carry obligations; individuals cannot easily check compliance"),
        ("ACCC Digital Platforms Inquiry",
         "found consumers have little practical visibility of who receives their data"),
        ("Nobody can check their own exposure",
         "the tooling exists for researchers, not for the person whose data it is"),
        ("Self-initiated individual project",
         "no external client; assessor is the approving stakeholder"),
    ], y=Inches(2.0), size=19, gap=Inches(1.05))
    notes(s, """
[1:25-2:25]
Why Australia specifically. Australian Privacy Principle 8 places obligations on
organisations that send personal information overseas. The ACCC's Digital
Platforms Inquiry found consumers have almost no practical visibility of this.

So there is a law about cross-border disclosure, and no way for an ordinary
person to see whether it is happening to them. That gap is what I'm building for.

This is my own initiative - there's no external client. The assessor is the
approving stakeholder.
""")

    # ---------------- 4. What I am building ----------------
    s = blank(prs)
    slide_title(s, "Two things you can install", kicker="scope")
    textbox(s, Inches(0.9), Inches(2.0), Inches(5.6), Inches(0.5),
            "1 · Chrome extension", size=24, color=BLUE, bold=True)
    textbox(s, Inches(0.9), Inches(2.6), Inches(5.6), Inches(2.6),
            ["Shows the problem on one machine",
             "",
             "• names the company behind every request",
             "• scores the page 0–100 for privacy risk",
             "• hashes your identifier locally and",
             "   watches for it leaving"],
            size=16, color=DARK, spacing=1.25)
    textbox(s, Inches(7.0), Inches(2.0), Inches(5.6), Inches(0.5),
            "2 · Raspberry Pi application", size=24, color=TEAL, bold=True)
    textbox(s, Inches(7.0), Inches(2.6), Inches(5.6), Inches(2.6),
            ["Acts on it for the whole network",
             "",
             "• adds owner and country attribution",
             "• blocks tracker domains by DNS for",
             "   every device, including phones and TVs",
             "• keeps history the browser forgets"],
            size=16, color=DARK, spacing=1.25)
    textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.9),
            "The extension shows you the problem. The Pi does something about it.",
            size=20, color=AMBER, bold=True)
    notes(s, """
[2:25-3:25]
The project is two things a person actually installs.

First, a Chrome extension. It watches every request the browser makes, names the
company behind each destination, scores the page for risk, and detects your own
identifier leaving as a hash.

Second, a Raspberry Pi application. It collects what the extension sees, adds
country attribution the browser can't do, and blocks tracker domains by DNS -
which covers every device on the network, including the phone and the TV that
can never run an extension.

One line to remember: the extension shows you the problem, the Pi does something
about it.
""")

    # ---------------- 5. Architecture ----------------
    s = blank(prs)
    slide_title(s, "How the pieces fit", kicker="architecture")
    diagram = ASSETS / "architecture_block_diagram.png"
    if diagram.exists():
        pic = s.shapes.add_picture(str(diagram), Inches(0.45), Inches(1.45),
                                   width=Inches(12.45))
        if pic.height > Inches(5.6):
            ratio = Inches(5.6) / pic.height
            pic.height = Inches(5.6)
            pic.width = Emu(int(pic.width * ratio))
            pic.left = Emu(int((SLIDE_W - pic.width) / 2))
    notes(s, """
[3:25-4:40]
This is the whole system on one slide.

On the left, the home network. The laptop runs the extension; the phone and the
TV cannot. In the middle, the Raspberry Pi: the DNS sinkhole at the top answers
tracker queries with 0.0.0.0, the FastAPI app receives telemetry, classify and
store handle attribution and storage, and the dashboard presents it.

Point at the orange bar along the bottom. That is the privacy boundary. The raw
email and the full page URLs never cross it. Only hashes are compared and only
hostnames are sent - because a URL carries your search terms and session ids, and
this is supposed to be a privacy tool.

Walk the diagram left to right with your hand, then turn back to the audience.
""")

    # ---------------- 6. Requirements ----------------
    s = blank(prs)
    slide_title(s, "What it has to do", kicker="requirements")
    textbox(s, Inches(0.9), Inches(1.8), Inches(5.6), Inches(0.45),
            "High-level — what", size=20, color=NAVY, bold=True)
    textbox(s, Inches(0.9), Inches(2.35), Inches(5.6), Inches(3.4),
            ["HR-1  Who receives the data, and who owns them",
             "HR-2  Which country, and is it offshore",
             "HR-3  Detect the identifier leaving as a hash",
             "HR-4  Covert third-party vs intentional use",
             "HR-5  Block for every device, not one browser",
             "HR-6  Present it so a non-specialist can act",
             "HR-7  Measure across 50 Australian sites"],
            size=15, color=DARK, spacing=1.45)
    textbox(s, Inches(7.0), Inches(1.8), Inches(5.6), Inches(0.45),
            "Low-level — how", size=20, color=NAVY, bold=True)
    textbox(s, Inches(7.0), Inches(2.35), Inches(5.6), Inches(3.4),
            ["Manifest V3 service worker observes requests",
             "WebCrypto + local MD5, four algorithms",
             "Every normalisation: digits, +61, Gmail dots",
             "Hostnames only — never URLs, never raw values",
             "GeoLite2 lookup on the Pi",
             "dnslib UDP server, upstream 1.1.1.1",
             "Playwright crawler over a frozen site list"],
            size=15, color=DARK, spacing=1.45)
    notes(s, """
[4:40-5:40]
Seven high-level requirements on the left - what the system has to do. Seven
matching implementation decisions on the right - how.

I'll call out two. HR-4: distinguishing covert from intentional. If you type your
email into a site's own search box, that's you choosing to. If that site quietly
forwards the hash to Facebook, that's the thing worth alarming about. Early on I
treated both the same, and everything looked like a breach.

And the normalisation row. A tracker doesn't hash what you typed - it hashes what
its own pipeline normalised first. I'll come back to that.

In the plan each of these is traced through to how it gets verified.
""")

    # ---------------- 7. Goals ----------------
    s = blank(prs)
    slide_title(s, "Goals I can be held to", kicker="smart goals")
    bullets(s, [
        ("G1 · Attribute owner and country for ≥ 90% of destinations",
         "by week 10, verified by manual spot-check against Tracker Radar"),
        ("G2 · Match a hashed identifier across 4 algorithms and 4 forms",
         "by week 10, evidenced by a reproducible test"),
        ("G3 · Zero misclassification of covert vs intentional",
         "by week 11, across the manual test set"),
        ("G4 · Block ≥ 75,000 domains with zero false positives",
         "by week 11, across twelve named Australian and global sites"),
        ("G5 · A non-specialist can read the result unaided",
         "by week 12, validated by a think-aloud walk-through"),
        ("G6 · Report H1–H4 with stated statistical methods",
         "by week 13, significant or not"),
    ], y=Inches(1.85), size=17, gap=Inches(0.83))
    notes(s, """
[5:40-6:25]
Six goals, each measurable and dated, and together they cover every requirement.

The one I'd point at is G4: zero false positives. Not "few". At the DNS layer a
false positive doesn't inconvenience one browser - it takes a website off the air
for everyone in the house. That's the number that has to be zero.

G6 says "significant or not" deliberately. If sectors don't differ, I report that.
I'm not designing this to confirm an answer I already like.

Don't read all six. Name the count, then talk about G4 and G6.
""")

    # ---------------- 8. Methodology ----------------
    s = blank(prs)
    slide_title(s, "Incremental layers, and why", kicker="methodology")
    textbox(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(0.5),
            "Each layer is a complete working system on its own",
            size=19, color=NAVY, bold=True)
    textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.5),
            ["L1  Extension alone   →   L2  + Pi collector   →   "
             "L3  + network-wide DNS blocking   →   L4  + 50-site study"],
            size=17, color=BLUE, bold=True)
    bullets(s, [
        ("Not waterfall — the riskiest unknowns are technical and early",
         "whether Manifest V3 allows the observation; whether a blocklist breaks real sites"),
        ("Not formal Scrum — ceremonies coordinate a team; there is one student",
         "short increments and working software kept; meeting structure dropped"),
        ("It already paid for itself",
         "wikipedia.org and github.com were classified as trackers — caught before that layer "
         "ever reached the network"),
    ], y=Inches(3.9), size=17, gap=Inches(0.95))
    notes(s, """
[6:25-7:40]
Methodology, and this is where I want to be concrete rather than quote a textbook.

I build in vertical layers. Each one is a complete working system by itself, and
later layers add capability without making earlier ones worthless.

Why not waterfall: my riskiest unknowns are technical and they sit early. Waterfall
would surface them in an integration phase with no time left to react.

Why not formal Scrum: the ceremonies exist to coordinate a team. There's one of me.
I kept short increments and working software; I dropped the meetings.

And here's the evidence it was the right call. My DNS blocklist parser classified
wikipedia.org and github.com as trackers - because filter rules that block a single
path were being read as blocking the whole domain. At the DNS layer that takes
Wikipedia off the air for every device in the house. I caught it because I tested
that layer in isolation before connecting it to anything.

This story is worth telling properly - it's the strongest thing on the slide.
""")

    # ---------------- 9. Tools and standards ----------------
    s = blank(prs)
    slide_title(s, "Tools chosen for a reason", kicker="tools & standards")
    textbox(s, Inches(0.9), Inches(1.8), Inches(5.7), Inches(3.9),
            ["Manifest V3 — the only platform Chrome accepts",
             "Python + FastAPI — runs comfortably on a Pi 5",
             "DuckDB — analytical queries, single file, no server",
             "dnslib — blocking logic stays readable and testable",
             "Playwright — a real browser, so real trackers fire",
             "GeoLite2 — free, offline, widely cited"],
            size=16, color=DARK, spacing=1.55)
    textbox(s, Inches(7.0), Inches(1.8), Inches(5.6), Inches(0.45),
            "Standards", size=20, color=NAVY, bold=True)
    textbox(s, Inches(7.0), Inches(2.35), Inches(5.6), Inches(3.4),
            ["Privacy Act 1988 — APP 8",
             "Manifest V3; least-privilege permissions",
             "Privacy by design — hashing stays client-side",
             "PEP 8, type annotations, semantic versioning",
             "Passive-observation research ethics",
             "   public pages, own traffic, no logins, no evasion"],
            size=15, color=DARK, spacing=1.5)
    textbox(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.8),
            "Every data source is public — EasyPrivacy, Tracker Radar, GeoLite2 — so any "
            "classification I report can be independently audited.",
            size=15, color=AMBER, bold=True)
    notes(s, """
[7:40-8:25]
Quickly on tools, because the plan justifies each one in detail.

Two worth naming here. DuckDB, because my questions are aggregate ones - "which
company receives the most requests" - and that's what a columnar database is for,
in a single file with no server to administer.

And dnslib rather than configuring dnsmasq, because I want the blocking logic to be
code I can read and test, not settings buried in a config file.

On standards: everything I classify comes from a public list. That means a marker
can check any claim I make. That's deliberate.
""")

    # ---------------- 10. Evidence ----------------
    s = blank(prs)
    slide_title(s, "Already built and tested", kicker="progress")
    bullets(s, [
        ("Extension v3.3 — capture, risk scoring, local hashing, malware blocking",
         "Manifest V3, running in Chrome today"),
        ("MD5 implementation verified against 12 test vectors",
         "including UTF-8 input and the 55/56/57/64-byte padding boundaries"),
        ("Identifier normalisation fixed a real miss",
         "typed as \"0412 345 678\", matched when a tracker hashes \"61412345678\" — the old single-form "
         "version matched nothing"),
        ("DNS sinkhole: zero false positives across twelve sites",
         "blocked domains answer 0.0.0.0 and :: ; everything else forwarded"),
        ("Telemetry is idempotent — a retried delivery double-counts nothing",
         "tested by replaying the same payload"),
        ("Study pipeline: 10 government sites crawled, 543 requests classified",
         "remaining 40 sites scheduled for weeks 9–10"),
    ], y=Inches(1.8), size=16, gap=Inches(0.82))
    notes(s, """
[8:25-9:25]
This is the slide I'd most like you to take away, because none of it is a plan -
it's all been run.

The MD5 I wrote is verified against twelve reference vectors, including the padding
boundaries at 55, 56 and 57 bytes, which is where hand-written MD5 usually breaks.

The normalisation line is my favourite. If you type your phone number as "0412 345
678" and a tracker hashes "61412345678", the old version of my code matched nothing
and reported no leak - silently. The fix hashes every plausible form. I have a test
that proves the old one missed and the new one matches.

And zero false positives on the DNS blocking, which is the G4 number.

Say "none of this is a plan, it's all been run" and then pause.
""")

    # ---------------- 11. Timeline and risks ----------------
    s = blank(prs)
    slide_title(s, "What is left, and what could go wrong", kicker="timeline & risk")
    textbox(s, Inches(0.9), Inches(1.75), Inches(5.7), Inches(0.45),
            "Remaining schedule", size=19, color=NAVY, bold=True)
    textbox(s, Inches(0.9), Inches(2.3), Inches(5.7), Inches(3.4),
            ["W9–10   Full 50-site crawl (G1, G2)",
             "W11       Network deployment, false-positive testing (G3, G4)",
             "W12       Dashboard and usability walk-through (G5)",
             "W13       Hypothesis testing and report (G6)",
             "W14       Final presentation"],
            size=15, color=DARK, spacing=1.6)
    textbox(s, Inches(7.0), Inches(1.75), Inches(5.6), Inches(0.45),
            "Top risks", size=19, color=NAVY, bold=True)
    textbox(s, Inches(7.0), Inches(2.3), Inches(5.6), Inches(3.4),
            ["A false positive takes a site off the air for the house",
             "   → separate lists for blocking and labelling; allow-test first",
             "The Pi becomes a single point of DNS failure",
             "   → documented rollback; blank secondary is a stated trade-off",
             "Chrome DNS-over-HTTPS bypasses the sinkhole",
             "   → stated limitation; the extension still sees these requests"],
            size=14, color=DARK, spacing=1.5)
    notes(s, """
[9:25-9:55]
Briefly: weeks 9 and 10 are the full crawl, 11 is deployment and the false-positive
testing, 12 and 13 are the dashboard and the analysis, 14 is the final presentation.

On risk - the honest one is the third. If someone turns on Chrome's secure DNS,
their queries go straight to Google and my sinkhole never sees them. I can't stop
that. But the extension still observes those requests, which is precisely why the
project is two deliverables and not one.

Don't rush this slide just because it's near the end.
""")

    # ---------------- 12. Close ----------------
    s = blank(prs)
    band(s, NAVY, Inches(2.2), Inches(0))
    textbox(s, Inches(0.9), Inches(0.72), Inches(11.5), Inches(1.0),
            "A hash is not anonymous.", size=38, color=WHITE, bold=True)
    textbox(s, Inches(0.95), Inches(1.55), Inches(11.5), Inches(0.5),
            "This project proves it on your own traffic — then blocks it.",
            size=19, color=RGBColor(0xBD, 0xD6, 0xF5))
    textbox(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(2.2),
            ["Extension — shows the problem on one machine",
             "Raspberry Pi — acts on it for every device on the network",
             "Everything classified from public, auditable sources"],
            size=20, color=DARK, spacing=1.9)
    textbox(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.9),
            ["Questions?", "github.com/thanhdat23102005/DataExo"],
            size=20, color=NAVY, bold=True, spacing=1.35)
    notes(s, """
[9:55-10:00]
To close: a hash is not anonymous, and this project proves it on your own traffic
rather than describing it. The extension shows you the problem; the Pi does
something about it for every device you own.

Happy to take questions.

Likely questions and short answers:
- "Why not just use Pi-hole?" Pi-hole blocks. It doesn't tell you who owns a
  destination, which country it's in, or whether your identifier went with it.
- "Is this legal?" Public pages only, my own traffic, my own identifier, no logins
  and no attacks - the same basis as published academic measurement work.
- "Why only 50 sites?" Because I verify every classification by hand. Fifty is a
  number I can audit; five thousand is a number I could only assert.
- "What if you find nothing?" I report that. It's a measurement, not an argument.
""")

    out = Path(__file__).parent / "Early_Stage_Presentation_DataExodus.pptx"
    prs.save(str(out))
    print(f"[done] {out.name}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
